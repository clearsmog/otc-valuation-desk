"""Market-data fetchers with an on-disk snapshot.

Every fetch writes the raw response under ``data/raw/`` so a valuation date can
be rerun offline and the tests never touch the network.

Sources (all free, no API key):
- Bank of England daily UK OIS spot curve (continuously compounded zero rates)
- Bank of England IADB: SONIA fixings (series IUDSOIA)
- Federal Reserve Bank of New York: SOFR fixings
- US Treasury daily par yield curve (used as a USD OIS proxy, see curves.py)
- ECB reference exchange rates (EUR/USD, EUR/GBP -> GBP/USD cross)
"""

from __future__ import annotations

import io
import zipfile
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pandas as pd
import requests

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "data" / "raw"
UA = {"User-Agent": "Mozilla/5.0 (otc-valuation-desk research)"}

BOE_OIS_ZIP = "https://www.bankofengland.co.uk/-/media/boe/files/statistics/yield-curves/latest-yield-curve-data.zip"
BOE_IADB = (
    "https://www.bankofengland.co.uk/boeapps/database/_iadb-fromshowcolumns.asp"
    "?csv.x=yes&Datefrom={start}&Dateto={end}&SeriesCodes=IUDSOIA&CSVF=TN&UsingCodes=Y&VPD=Y&VFD=N"
)
NYFED_SOFR = "https://markets.newyorkfed.org/api/rates/secured/sofr/search.json?startDate={start}&endDate={end}"
TREASURY = (
    "https://home.treasury.gov/resource-center/data-chart-center/interest-rates/daily-treasury-rates.csv/"
    "{year}/all?type=daily_treasury_yield_curve&field_tdr_date_value={year}&page&_format=csv"
)
ECB_FX = "https://data-api.ecb.europa.eu/service/data/EXR/D.USD+GBP.EUR.SP00.A?startPeriod={start}&endPeriod={end}&format=csvdata"


def _get(url: str, path: Path, refresh: bool) -> bytes:
    if path.exists() and not refresh:
        return path.read_bytes()
    r = requests.get(url, headers=UA, timeout=60)
    r.raise_for_status()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(r.content)
    return r.content


# --- Bank of England OIS curve ---------------------------------------------


def _parse_boe_sheet(ws) -> pd.DataFrame:
    rows = list(ws.iter_rows(values_only=True))
    years_row = next(r for r in rows if r and r[0] == "years:")
    tenors = [float(x) for x in years_row[1:] if x is not None]
    data = {}
    for r in rows:
        if isinstance(r[0], datetime):
            vals = [r[i + 1] for i in range(len(tenors))]
            data[r[0].date()] = [
                float(v) if isinstance(v, (int, float)) else float("nan") for v in vals
            ]
    return pd.DataFrame.from_dict(data, orient="index", columns=tenors).sort_index()


def boe_ois_curve(month: str, refresh: bool = False) -> pd.DataFrame:
    """UK OIS spot curve: index = curve date, columns = maturity in years,
    values = continuously compounded zero rates in percent.

    Joins the short-end sheet (1-60 months) with the long-end sheet (0.5y
    steps); the short end wins where both have a point. ``month`` ("YYYY-MM")
    names the cached copy: the Bank replaces the "latest" file every month, so
    the cache key keeps a past valuation date reproducible.
    """
    raw = _get(BOE_OIS_ZIP, RAW / "boe" / f"ois_{month}.zip", refresh)
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        name = next(n for n in z.namelist() if n.startswith("OIS daily"))
        wb = openpyxl.load_workbook(io.BytesIO(z.read(name)), data_only=True)
    short = _parse_boe_sheet(wb["3. spot, short end"])
    long = _parse_boe_sheet(wb["4. spot curve"])
    long = long[[c for c in long.columns if c > short.columns.max() + 1e-9]]
    out = short.join(long, how="outer")
    out.columns = [round(c, 6) for c in out.columns]
    return out.dropna(axis=1, how="all")


# --- Fixings ------------------------------------------------------------------


def sonia_fixings(start: date, end: date, refresh: bool = False) -> pd.Series:
    """SONIA in percent, indexed by the date the rate applies to."""
    url = BOE_IADB.format(
        start=start.strftime("%d/%b/%Y"), end=end.strftime("%d/%b/%Y")
    )
    raw = _get(url, RAW / "boe" / f"sonia_{start}_{end}.csv", refresh)
    df = pd.read_csv(io.BytesIO(raw))
    s = pd.Series(
        df["IUDSOIA"].astype(float).values,
        index=pd.to_datetime(df["DATE"], format="%d %b %Y").dt.date,
    )
    return s.sort_index().rename("SONIA")


def sofr_fixings(start: date, end: date, refresh: bool = False) -> pd.Series:
    """SOFR in percent, indexed by effective date."""
    raw = _get(
        NYFED_SOFR.format(start=start, end=end),
        RAW / "nyfed" / f"sofr_{start}_{end}.json",
        refresh,
    )
    recs = pd.read_json(io.BytesIO(raw))["refRates"]
    s = pd.Series(
        {date.fromisoformat(r["effectiveDate"]): float(r["percentRate"]) for r in recs}
    )
    return s.sort_index().rename("SOFR")


# --- US Treasury par curve ------------------------------------------------------

TREASURY_TENORS = {
    "1 Mo": 1 / 12,
    "2 Mo": 2 / 12,
    "3 Mo": 0.25,
    "4 Mo": 4 / 12,
    "6 Mo": 0.5,
    "1 Yr": 1.0,
    "2 Yr": 2.0,
    "3 Yr": 3.0,
    "5 Yr": 5.0,
    "7 Yr": 7.0,
    "10 Yr": 10.0,
    "20 Yr": 20.0,
    "30 Yr": 30.0,
}


def treasury_par(year: int, refresh: bool = False) -> pd.DataFrame:
    """Par yields in percent: index = date, columns = maturity in years."""
    raw = _get(
        TREASURY.format(year=year), RAW / "treasury" / f"par_{year}.csv", refresh
    )
    df = pd.read_csv(io.BytesIO(raw))
    df.index = pd.to_datetime(df["Date"], format="%m/%d/%Y").dt.date
    df = df[[c for c in TREASURY_TENORS if c in df.columns]].astype(float)
    df.columns = [TREASURY_TENORS[c] for c in df.columns]
    return df.sort_index()


# --- ECB FX -------------------------------------------------------------------


def ecb_fx(start: date, end: date, refresh: bool = False) -> pd.DataFrame:
    """ECB reference rates EUR/USD and EUR/GBP, plus the GBP/USD cross."""
    raw = _get(
        ECB_FX.format(start=start, end=end),
        RAW / "ecb" / f"fx_{start}_{end}.csv",
        refresh,
    )
    df = pd.read_csv(io.BytesIO(raw), usecols=["CURRENCY", "TIME_PERIOD", "OBS_VALUE"])
    wide = df.pivot(index="TIME_PERIOD", columns="CURRENCY", values="OBS_VALUE")
    wide.index = [date.fromisoformat(d) for d in wide.index]
    out = pd.DataFrame({"EURUSD": wide["USD"], "EURGBP": wide["GBP"]})
    out["GBPUSD"] = out["EURUSD"] / out["EURGBP"]
    return out.sort_index()


BOE_OIS_ARCHIVE = "https://www.bankofengland.co.uk/-/media/boe/files/statistics/yield-curves/oisddata.zip"


def boe_ois_archive(refresh: bool = False) -> pd.DataFrame:
    """UK OIS spot curve history from 2025 (the Bank's archive, updated about
    monthly). Only the 2025-onwards rows are cached, as CSV, to keep the
    repository small."""
    path = RAW / "boe" / "ois_archive_2025_on.csv"
    if path.exists() and not refresh:
        df = pd.read_csv(path, index_col=0)
        df.index = [date.fromisoformat(d) for d in df.index]
        df.columns = [float(c) for c in df.columns]
        return df
    r = requests.get(BOE_OIS_ARCHIVE, headers=UA, timeout=120)
    r.raise_for_status()
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        name = next(n for n in z.namelist() if "2025 to present" in n)
        wb = openpyxl.load_workbook(io.BytesIO(z.read(name)), data_only=True)
    short = _parse_boe_sheet(wb["3. spot, short end"])
    long = _parse_boe_sheet(wb["4. spot curve"])
    long = long[[c for c in long.columns if c > short.columns.max() + 1e-9]]
    out = short.join(long, how="outer")
    out.columns = [round(c, 6) for c in out.columns]
    out = out.dropna(axis=1, how="all")
    path.parent.mkdir(parents=True, exist_ok=True)
    out.to_csv(path)
    return out


def boe_curve_row(asof: date) -> pd.Series:
    """The BoE OIS zero curve for one date: the archive if it has the date,
    otherwise the current-month file."""
    arch = boe_ois_archive()
    if asof in arch.index:
        return arch.loc[asof].dropna()
    return boe_ois_curve(asof.strftime("%Y-%m")).loc[asof].dropna()
