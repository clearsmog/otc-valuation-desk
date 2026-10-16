"""The daily valuation run.

    uv run python -m otc_valuation_desk.cycle 2026-10-15 --prior 2026-10-14

Steps: market-data QC -> curves -> own valuation and sensitivities ->
QuantLib reference check -> dealer marks -> price challenges -> report
(``outputs/<asof>/valuation_report.md`` and ``.xlsx``).
"""

from __future__ import annotations

import argparse
from dataclasses import asdict
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import challenge as CH
from . import data as D
from . import marks as M
from . import pricing as P
from . import qc
from . import reference as R
from .curves import t_of
from .market import Market, build_market
from .trades import FXOption, OISSwap, Swaption, Trade, load_portfolio

ROOT = Path(__file__).resolve().parents[2]
PORTFOLIO = ROOT / "data" / "trades" / "client_portfolio.csv"
REF_TOL = {"OIS": 0.1, "FXFWD": 0.5, "FXOPT": 0.05, "SWPTN": 0.05, "XCCY": 0.5}  # bp of PV01, bp of spot, vol pts, bp normal vol


def run_qc(mkt: Market) -> pd.DataFrame:
    asof = mkt.asof
    lookback = asof - timedelta(days=60)
    ex: list[qc.QCException] = []
    for name, cal in (("SONIA", "GBP"), ("SOFR", "USD")):
        s = mkt.fixings[name]
        s = s[(s.index >= lookback) & (s.index < asof)]
        ex += qc.check_missing(s, name, lookback, asof - timedelta(days=1), cal)
        ex += qc.check_stale(s, name, stale_days=5)
        ex += qc.check_jumps(s, name, threshold=10.0, unit="bp")
    fx = D.ecb_fx(date(2026, 9, 1), date(2026, 10, 16))["GBPUSD"]
    ex += qc.check_jumps(fx[fx.index <= asof], "ECB", threshold=1.5, unit="pct")
    ust = D.treasury_par(asof.year)
    ust = ust[(ust.index >= lookback) & (ust.index <= asof)]
    for col in ust.columns:
        ex += qc.check_jumps(
            ust[col].rename(f"UST {col:g}y"), "Treasury", threshold=15.0, unit="bp"
        )
    boe = D.boe_ois_curve(asof.strftime("%Y-%m")).loc[asof]
    ex += qc.check_zero_curve(boe.index.values, boe.values, "BoE OIS", asof)
    return qc.to_frame(ex)


def curve_table(mkt: Market, qm: R.QLMarket) -> pd.DataFrame:
    rows = []
    for ccy, c in mkt.curves.items():
        for t in (0.25, 1, 2, 5, 10, 20):
            d = mkt.asof + timedelta(days=round(t * 365))
            rows.append(
                {
                    "curve": c.label,
                    "tenor_y": t,
                    "zero_pct": round(c.zero(t_of(mkt.asof, d)) * 100, 4),
                    "df_own": c.df_at(d),
                    "df_quantlib": qm.df(ccy, d),
                    "df_diff_e6": (c.df_at(d) - qm.df(ccy, d)) * 1e6,
                }
            )
    return pd.DataFrame(rows)


def _ref_unit(trade: Trade, sens: dict[str, float]) -> float:
    if isinstance(trade, OISSwap):
        return abs(sens[f"pv01_{trade.ccy}"])
    if isinstance(trade, FXOption):
        return abs(sens["vega_1pt"])
    if isinstance(trade, Swaption):
        return abs(sens["vega_1bp"])
    return abs(sens["fx_delta_1pct"]) / 100.0


def valuation_table(trades: list[Trade], mkt: Market, qm: R.QLMarket) -> pd.DataFrame:
    rows = []
    for t in trades:
        own, ref = P.price(t, mkt), R.ql_price(t, qm)
        sens = P.sensitivities(t, mkt)
        unit = _ref_unit(t, sens)
        diff_units = (own - ref) / unit if unit else 0.0
        ccy = P.pv_ccy(t)
        rows.append(
            {
                "trade_id": t.trade_id,
                "product": t.product,
                "counterparty": t.counterparty,
                "ccy": ccy,
                "pv_own": own,
                "pv_quantlib": ref,
                "own_minus_ql": own - ref,
                "diff_in_risk_units": diff_units,
                "ref_check": "OK"
                if abs(diff_units) <= REF_TOL[t.product]
                else "MODEL CHECK",
                "pv_usd": P.to_usd(own, ccy, mkt),
                **sens,
            }
        )
    return pd.DataFrame(rows)


# --- report writers --------------------------------------------------------------


def _md_table(df: pd.DataFrame, floatfmt: str = ",.0f") -> str:
    def fmt(v):
        if isinstance(v, float):
            return format(v, floatfmt)
        return str(v)

    head = "| " + " | ".join(df.columns) + " |\n|" + "---|" * len(df.columns) + "\n"
    return head + "\n".join(
        "| " + " | ".join(fmt(v) for v in r) + " |" for r in df.itertuples(index=False)
    )


def write_markdown(path: Path, mkt: Market, qc_df, curves_df, val_df, ch_df) -> None:
    counts = ch_df["status"].value_counts().to_dict()
    lines = [
        f"# Valuation report — {mkt.asof}",
        "",
        "Fictional client portfolio (20 trades). GBP curve: Bank of England OIS. "
        "USD curve: US Treasury par yields as a SOFR OIS **proxy**. FX spot: ECB reference rate. "
        "FX vols and dealer marks are **synthetic**.",
        "",
        "## Summary",
        "",
        f"- Portfolio PV (USD, own model): {val_df['pv_usd'].sum():,.0f}",
        f"- QuantLib reference: {(val_df['ref_check'] == 'OK').sum()}/{len(val_df)} trades within tolerance",
        f"- Market-data QC exceptions: {len(qc_df)} ({(qc_df['severity'] == 'FAIL').sum() if len(qc_df) else 0} FAIL)",
        "- Price challenges: "
        + ", ".join(f"{k} {v}" for k, v in sorted(counts.items())),
        "",
        "## Market-data QC",
        "",
        _md_table(qc_df) if len(qc_df) else "No exceptions.",
        "",
        "## Curves (own bootstrap vs QuantLib)",
        "",
        _md_table(curves_df, ",.6f"),
        "",
        "## Valuations",
        "",
        _md_table(
            val_df[
                [
                    "trade_id",
                    "product",
                    "counterparty",
                    "ccy",
                    "pv_own",
                    "pv_quantlib",
                    "own_minus_ql",
                    "diff_in_risk_units",
                    "ref_check",
                ]
            ]
        ),
        "",
        "## Price challenges (dealer mark minus our value)",
        "",
        _md_table(
            ch_df[
                [
                    "trade_id",
                    "counterparty",
                    "ccy",
                    "our_value",
                    "dealer_mark",
                    "difference",
                    "tolerance",
                    "status",
                    "cause",
                    "detail",
                    "action",
                ]
            ]
        ),
        "",
    ]
    path.write_text("\n".join(lines))


FILLS = {
    "WITHIN_TOLERANCE": "E2EFDA",
    "EXPLAINED": "FFF2CC",
    "LIKELY": "FCE4D6",
    "UNEXPLAINED": "F8CBAD",
    "MODEL CHECK": "F8CBAD",
    "FAIL": "F8CBAD",
    "WARN": "FFF2CC",
}


def write_excel(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for i, col in enumerate(df.columns, start=1):
                width = (
                    max(
                        len(str(col)),
                        *(
                            len(f"{v:,.0f}" if isinstance(v, float) else str(v))
                            for v in df[col]
                        ),
                    )
                    if len(df)
                    else len(col)
                )
                ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)
                if pd.api.types.is_float_dtype(df[col]):
                    fmt = (
                        "0.000000"
                        if col.startswith("df_")
                        else "#,##0.00"
                        if "units" in col
                        else "#,##0"
                    )
                    for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                        row[0].number_format = fmt
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str) and cell.value in FILLS:
                        cell.fill = PatternFill("solid", fgColor=FILLS[cell.value])


def run_daily(
    asof: date, prior: date, outdir: Path | None = None
) -> dict[str, pd.DataFrame]:
    outdir = outdir or ROOT / "outputs" / str(asof)
    outdir.mkdir(parents=True, exist_ok=True)
    mkt = build_market(asof)
    qm = R.QLMarket(mkt)
    trades = load_portfolio(PORTFOLIO)

    qc_df = run_qc(mkt)
    curves_df = curve_table(mkt, qm)
    val_df = valuation_table(trades, mkt, qm)
    marks, key = M.generate_marks(trades, mkt, prior)
    pd.DataFrame([asdict(m) for m in marks]).to_csv(
        outdir / "dealer_marks.csv", index=False
    )
    pd.DataFrame(key).to_csv(outdir / "answer_key_planted_causes.csv", index=False)
    ch_df = pd.DataFrame(
        [asdict(r) for r in CH.run_challenges(trades, marks, mkt, prior)]
    )

    sheets = {
        "QC": qc_df,
        "Curves": curves_df,
        "Valuations": val_df,
        "Challenges": ch_df,
    }
    write_markdown(outdir / "valuation_report.md", mkt, qc_df, curves_df, val_df, ch_df)
    write_excel(outdir / "valuation_report.xlsx", sheets)
    return sheets


def main() -> None:
    ap = argparse.ArgumentParser(description="Daily OTC valuation run")
    ap.add_argument("asof", type=date.fromisoformat)
    ap.add_argument(
        "--prior",
        type=date.fromisoformat,
        required=True,
        help="previous business day (for dealer replay)",
    )
    a = ap.parse_args()
    sheets = run_daily(a.asof, a.prior)
    ch = sheets["Challenges"]
    print(
        f"{a.asof}: {len(sheets['Valuations'])} trades, {len(sheets['QC'])} QC exceptions"
    )
    print(ch["status"].value_counts().to_string())
    print(f"report: outputs/{a.asof}/valuation_report.md, .xlsx")


if __name__ == "__main__":
    main()
